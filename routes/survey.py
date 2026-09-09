"""학교관리 > 설문조사.

제목과 문항을 만들어 두고, 수신자 휴대폰번호 목록에 SOLAPI 문자로 응답 URL을
보내 결과를 모으는 메뉴. 응답자는 인트라넷 계정 없이 링크만으로 참여한다.
"""

from __future__ import annotations

import io
import json
import re
import secrets
from datetime import datetime

from flask import (
    Blueprint, jsonify, render_template, request, send_file, session, url_for,
)

from .database import get_db
from .solapi_settings import (
    DEFAULT_SENDER_NAME,
    SMS_BYTE_LIMIT,
    format_phone,
    get_settings as get_solapi_settings,
    mask_phone,
    message_byte_length,
    normalize_phone,
    resolve_public_origin,
    send_bulk_text,
)

survey_bp = Blueprint('survey', __name__)

QUESTION_TYPES = {'single', 'multiple', 'text', 'scale'}
QUESTION_TYPE_LABELS = {
    'single': '단일선택',
    'multiple': '복수선택',
    'text': '주관식',
    'scale': '5점 척도',
}
SURVEY_STATUSES = {'draft': '작성중', 'open': '진행중', 'closed': '마감'}
SCALE_DEFAULT_OPTIONS = ['매우 그렇다', '그렇다', '보통이다', '아니다', '전혀 아니다']
DEFAULT_SMS_TEMPLATE = '[{기관명}] {제목} 설문에 참여해 주세요.\n{링크}'
RECIPIENT_COLUMNS = ('번호', '구분', '이름', '핸드폰번호', '비고')
RECIPIENT_SAMPLE_ROWS = (
    (1, '신풍초', '홍길동', '010-1234-5678', '3학년 담당'),
    (2, '신풍초', '김새담', '010-2345-6789', ''),
    (3, '본사', '이하늘', '010-3456-7890', '주말 연락 요망'),
)
MAX_QUESTIONS = 50
MAX_OPTIONS = 20
MAX_RECIPIENTS = 2000


# ---------------------------------------------------------------------------
# 스키마
# ---------------------------------------------------------------------------
def init_survey_schema(conn=None):
    owns_connection = conn is None
    if owns_connection:
        conn = get_db()
    try:
        conn.executescript('''
            CREATE TABLE IF NOT EXISTS surveys (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                description TEXT,
                status TEXT NOT NULL DEFAULT 'draft',
                public_token TEXT NOT NULL UNIQUE,
                starts_on TEXT,
                ends_on TEXT,
                allow_public_link INTEGER NOT NULL DEFAULT 0,
                allow_anonymous INTEGER NOT NULL DEFAULT 0,
                sms_template TEXT,
                created_by TEXT,
                created_by_name TEXT,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS survey_questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                survey_id INTEGER NOT NULL,
                sort_order INTEGER NOT NULL DEFAULT 0,
                qtype TEXT NOT NULL DEFAULT 'single',
                title TEXT NOT NULL,
                options TEXT,
                is_required INTEGER NOT NULL DEFAULT 1,
                FOREIGN KEY (survey_id) REFERENCES surveys(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_survey_questions_survey
                ON survey_questions(survey_id, sort_order);

            CREATE TABLE IF NOT EXISTS survey_recipients (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                survey_id INTEGER NOT NULL,
                name TEXT,
                phone TEXT NOT NULL,
                category TEXT,
                note TEXT,
                token TEXT NOT NULL UNIQUE,
                send_status TEXT NOT NULL DEFAULT 'ready',
                send_error TEXT,
                message_id TEXT,
                sent_at DATETIME,
                responded_at DATETIME,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                UNIQUE (survey_id, phone),
                FOREIGN KEY (survey_id) REFERENCES surveys(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_survey_recipients_survey
                ON survey_recipients(survey_id);

            CREATE TABLE IF NOT EXISTS survey_responses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                survey_id INTEGER NOT NULL,
                recipient_id INTEGER,
                is_anonymous INTEGER NOT NULL DEFAULT 0,
                submitted_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (survey_id) REFERENCES surveys(id) ON DELETE CASCADE,
                FOREIGN KEY (recipient_id) REFERENCES survey_recipients(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_survey_responses_survey
                ON survey_responses(survey_id);

            CREATE TABLE IF NOT EXISTS survey_answers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                response_id INTEGER NOT NULL,
                question_id INTEGER NOT NULL,
                answer_text TEXT,
                answer_options TEXT,
                FOREIGN KEY (response_id) REFERENCES survey_responses(id) ON DELETE CASCADE,
                FOREIGN KEY (question_id) REFERENCES survey_questions(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_survey_answers_response
                ON survey_answers(response_id);
            CREATE INDEX IF NOT EXISTS idx_survey_answers_question
                ON survey_answers(question_id);

            CREATE TABLE IF NOT EXISTS survey_send_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                survey_id INTEGER NOT NULL,
                sent_by TEXT,
                total INTEGER NOT NULL DEFAULT 0,
                success INTEGER NOT NULL DEFAULT 0,
                failed INTEGER NOT NULL DEFAULT 0,
                memo TEXT,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (survey_id) REFERENCES surveys(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_survey_send_logs_survey
                ON survey_send_logs(survey_id, created_at DESC);
        ''')
        _add_missing_columns(conn)
        conn.commit()
    finally:
        if owns_connection:
            conn.close()


def _add_missing_columns(conn):
    """먼저 만들어진 설치본에도 뒤에 추가한 열을 한 번씩 붙여 준다."""
    additions = (
        ('surveys', 'allow_anonymous', 'INTEGER NOT NULL DEFAULT 0'),
        ('survey_recipients', 'category', 'TEXT'),
        ('survey_recipients', 'note', 'TEXT'),
        ('survey_responses', 'is_anonymous', 'INTEGER NOT NULL DEFAULT 0'),
    )
    for table, column, definition in additions:
        columns = {
            row[1] for row in conn.execute(f'PRAGMA table_info({table})').fetchall()
        }
        if column not in columns:
            conn.execute(f'ALTER TABLE {table} ADD COLUMN {column} {definition}')


# ---------------------------------------------------------------------------
# 공통 도우미
# ---------------------------------------------------------------------------
def _actor():
    return {
        'emp_no': str(session.get('emp_no') or ''),
        'name': str(session.get('user_name') or ''),
    }


def _is_admin():
    if str(session.get('user_name') or '') == 'admin':
        return True
    try:
        return int(session.get('user_level', 99)) <= 2
    except (TypeError, ValueError):
        return False


def _can_manage(row):
    """설문을 수정·삭제할 수 있는지 확인한다. 등록자 본인과 관리자만 허용."""
    if _is_admin():
        return True
    emp_no = str(session.get('emp_no') or '')
    return bool(emp_no) and str(row['created_by'] or '') == emp_no


def _new_token():
    return secrets.token_urlsafe(24)


def _json_error(message, code=400):
    return jsonify({'status': 'error', 'message': message}), code


def _load_options(raw):
    try:
        parsed = json.loads(raw or '[]')
    except (TypeError, ValueError):
        return []
    return [str(item) for item in parsed] if isinstance(parsed, list) else []


def _clean_date(value):
    text = str(value or '').strip()
    if not text:
        return ''
    try:
        return datetime.strptime(text, '%Y-%m-%d').strftime('%Y-%m-%d')
    except ValueError:
        raise ValueError('날짜는 YYYY-MM-DD 형식으로 입력해 주세요.')


def _normalize_questions(payload):
    """화면에서 넘어온 문항 목록을 저장 가능한 형태로 검증한다."""
    if not isinstance(payload, list) or not payload:
        raise ValueError('설문 문항을 1개 이상 추가해 주세요.')
    if len(payload) > MAX_QUESTIONS:
        raise ValueError(f'설문 문항은 최대 {MAX_QUESTIONS}개까지 만들 수 있습니다.')

    questions = []
    for index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            raise ValueError(f'{index}번 문항 형식을 확인해 주세요.')
        title = str(item.get('title') or '').strip()
        if not title:
            raise ValueError(f'{index}번 문항의 질문 내용을 입력해 주세요.')
        if len(title) > 300:
            raise ValueError(f'{index}번 문항은 300자 이내로 입력해 주세요.')
        qtype = str(item.get('qtype') or 'single').strip()
        if qtype not in QUESTION_TYPES:
            raise ValueError(f'{index}번 문항의 유형을 확인해 주세요.')

        options = []
        if qtype in {'single', 'multiple'}:
            raw_options = item.get('options')
            if isinstance(raw_options, list):
                options = [str(opt).strip() for opt in raw_options if str(opt).strip()]
            if len(options) < 2:
                raise ValueError(f'{index}번 문항의 보기를 2개 이상 입력해 주세요.')
            if len(options) > MAX_OPTIONS:
                raise ValueError(f'{index}번 문항의 보기는 최대 {MAX_OPTIONS}개입니다.')
            if len(set(options)) != len(options):
                raise ValueError(f'{index}번 문항에 같은 보기가 중복되었습니다.')
        elif qtype == 'scale':
            options = list(SCALE_DEFAULT_OPTIONS)

        questions.append({
            'sort_order': index,
            'qtype': qtype,
            'title': title,
            'options': json.dumps(options, ensure_ascii=False),
            'is_required': 0 if str(item.get('is_required')) in {'0', 'False', 'false'} else 1,
        })
    return questions


def _survey_row(conn, survey_id):
    return conn.execute('SELECT * FROM surveys WHERE id=?', (survey_id,)).fetchone()


def _questions_of(conn, survey_id):
    rows = conn.execute(
        'SELECT * FROM survey_questions WHERE survey_id=? ORDER BY sort_order, id',
        (survey_id,),
    ).fetchall()
    return [{
        'id': row['id'],
        'sort_order': row['sort_order'],
        'qtype': row['qtype'],
        'qtype_label': QUESTION_TYPE_LABELS.get(row['qtype'], row['qtype']),
        'title': row['title'],
        'options': _load_options(row['options']),
        'is_required': bool(row['is_required']),
    } for row in rows]


def _survey_dict(row, *, response_count=0, recipient_count=0, sent_count=0):
    return {
        'id': row['id'],
        'title': row['title'],
        'description': row['description'] or '',
        'status': row['status'],
        'status_label': SURVEY_STATUSES.get(row['status'], row['status']),
        'public_token': row['public_token'],
        'starts_on': row['starts_on'] or '',
        'ends_on': row['ends_on'] or '',
        'allow_public_link': bool(row['allow_public_link']),
        'allow_anonymous': bool(row['allow_anonymous']),
        'sms_template': row['sms_template'] or DEFAULT_SMS_TEMPLATE,
        'created_by': row['created_by'] or '',
        'created_by_name': row['created_by_name'] or '',
        'created_at': str(row['created_at'] or '')[:16],
        'updated_at': str(row['updated_at'] or '')[:16],
        'response_count': response_count,
        'recipient_count': recipient_count,
        'sent_count': sent_count,
        'can_manage': _can_manage(row),
    }


def _public_link(token, origin=None):
    base = origin if origin is not None else resolve_public_origin()
    return f"{base}{url_for('survey.respond', token=token)}"


def _survey_is_open(row):
    """응답을 받을 수 있는 기간·상태인지 확인하고 안내 문구를 돌려준다."""
    if row['status'] == 'draft':
        return False, '아직 시작되지 않은 설문입니다.'
    if row['status'] == 'closed':
        return False, '마감된 설문입니다. 참여해 주셔서 감사합니다.'
    today = datetime.now().strftime('%Y-%m-%d')
    if row['starts_on'] and today < row['starts_on']:
        return False, f"{row['starts_on']}부터 참여할 수 있는 설문입니다."
    if row['ends_on'] and today > row['ends_on']:
        return False, f"{row['ends_on']}에 마감된 설문입니다."
    return True, ''


def _render_sms(template, *, title, link, sender_name):
    text = str(template or DEFAULT_SMS_TEMPLATE)
    return (
        text.replace('{기관명}', sender_name)
            .replace('{제목}', title)
            .replace('{링크}', link)
            .strip()
    )


# ---------------------------------------------------------------------------
# 화면
# ---------------------------------------------------------------------------
@survey_bp.route('/')
def index():
    return render_template('survey/index.html')


@survey_bp.route('/new')
def new_survey():
    return render_template('survey/editor.html', survey_id=0)


@survey_bp.route('/<int:survey_id>/edit')
def edit_survey(survey_id):
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return '설문을 찾을 수 없습니다.', 404
        if not _can_manage(row):
            return '이 설문을 수정할 권한이 없습니다.', 403
    finally:
        conn.close()
    return render_template('survey/editor.html', survey_id=survey_id)


@survey_bp.route('/<int:survey_id>/result')
def result_page(survey_id):
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return '설문을 찾을 수 없습니다.', 404
    finally:
        conn.close()
    return render_template('survey/result.html', survey_id=survey_id)


# ---------------------------------------------------------------------------
# 관리 API
# ---------------------------------------------------------------------------
@survey_bp.route('/api/list')
def api_list():
    keyword = str(request.args.get('q') or '').strip()
    status = str(request.args.get('status') or '').strip()
    conn = get_db()
    try:
        sql = '''
            SELECT s.*,
                   (SELECT COUNT(*) FROM survey_responses r WHERE r.survey_id=s.id)
                       AS response_count,
                   (SELECT COUNT(*) FROM survey_recipients c WHERE c.survey_id=s.id)
                       AS recipient_count,
                   (SELECT COUNT(*) FROM survey_recipients c
                     WHERE c.survey_id=s.id AND c.send_status='sent') AS sent_count
              FROM surveys s
        '''
        conditions, params = [], []
        if keyword:
            conditions.append('(s.title LIKE ? OR s.description LIKE ?)')
            params.extend([f'%{keyword}%', f'%{keyword}%'])
        if status in SURVEY_STATUSES:
            conditions.append('s.status=?')
            params.append(status)
        if conditions:
            sql += ' WHERE ' + ' AND '.join(conditions)
        sql += ' ORDER BY s.id DESC LIMIT 300'
        rows = conn.execute(sql, params).fetchall()
        items = [
            _survey_dict(
                row,
                response_count=row['response_count'],
                recipient_count=row['recipient_count'],
                sent_count=row['sent_count'],
            )
            for row in rows
        ]
    finally:
        conn.close()
    return jsonify({
        'status': 'success',
        'items': items,
        'statuses': SURVEY_STATUSES,
        'sms_ready': bool(get_solapi_settings().get('sms_configured')),
    })


@survey_bp.route('/api/<int:survey_id>')
def api_detail(survey_id):
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        counts = conn.execute(
            '''SELECT
                 (SELECT COUNT(*) FROM survey_responses WHERE survey_id=?) AS responses,
                 (SELECT COUNT(*) FROM survey_recipients WHERE survey_id=?) AS recipients,
                 (SELECT COUNT(*) FROM survey_recipients
                   WHERE survey_id=? AND send_status='sent') AS sent''',
            (survey_id, survey_id, survey_id),
        ).fetchone()
        survey = _survey_dict(
            row,
            response_count=counts['responses'],
            recipient_count=counts['recipients'],
            sent_count=counts['sent'],
        )
        origin = resolve_public_origin()
        survey['questions'] = _questions_of(conn, survey_id)
        survey['public_link'] = _public_link(row['public_token'], origin)
        recipients = [{
            'id': item['id'],
            'name': item['name'] or '',
            'category': item['category'] or '',
            'note': item['note'] or '',
            'phone': format_phone(item['phone']),
            'phone_masked': mask_phone(item['phone']),
            'send_status': item['send_status'],
            'send_error': item['send_error'] or '',
            'sent_at': str(item['sent_at'] or '')[:16],
            'responded_at': str(item['responded_at'] or '')[:16],
            'link': _public_link(item['token'], origin),
        } for item in conn.execute(
            'SELECT * FROM survey_recipients WHERE survey_id=? ORDER BY id',
            (survey_id,),
        ).fetchall()]
        logs = [{
            'sent_by': item['sent_by'] or '',
            'total': item['total'],
            'success': item['success'],
            'failed': item['failed'],
            'memo': item['memo'] or '',
            'created_at': str(item['created_at'] or '')[:16],
        } for item in conn.execute(
            'SELECT * FROM survey_send_logs WHERE survey_id=? ORDER BY id DESC LIMIT 50',
            (survey_id,),
        ).fetchall()]
    finally:
        conn.close()
    solapi = get_solapi_settings()
    return jsonify({
        'status': 'success',
        'survey': survey,
        'recipients': recipients,
        'send_logs': logs,
        'sms_ready': bool(solapi.get('sms_configured')),
        'sender_name': solapi.get('sender_name') or DEFAULT_SENDER_NAME,
        'default_template': DEFAULT_SMS_TEMPLATE,
    })


@survey_bp.route('/api/save', methods=['POST'])
def api_save():
    payload = request.get_json(silent=True) or {}
    survey_id = payload.get('id')
    title = str(payload.get('title') or '').strip()
    if not title:
        return _json_error('설문 제목을 입력해 주세요.')
    if len(title) > 200:
        return _json_error('설문 제목은 200자 이내로 입력해 주세요.')
    description = str(payload.get('description') or '').strip()[:2000]
    template = str(payload.get('sms_template') or DEFAULT_SMS_TEMPLATE).strip()[:500]
    if '{링크}' not in template:
        return _json_error('문자 내용에는 설문 주소가 들어갈 {링크} 를 반드시 넣어 주세요.')
    allow_public_link = 1 if payload.get('allow_public_link') else 0
    allow_anonymous = 1 if payload.get('allow_anonymous') else 0

    try:
        starts_on = _clean_date(payload.get('starts_on'))
        ends_on = _clean_date(payload.get('ends_on'))
        questions = _normalize_questions(payload.get('questions'))
    except ValueError as exc:
        return _json_error(str(exc))
    if starts_on and ends_on and starts_on > ends_on:
        return _json_error('설문 종료일이 시작일보다 빠릅니다.')

    actor = _actor()
    conn = get_db()
    try:
        if survey_id:
            row = _survey_row(conn, survey_id)
            if not row:
                return _json_error('설문을 찾을 수 없습니다.', 404)
            if not _can_manage(row):
                return _json_error('이 설문을 수정할 권한이 없습니다.', 403)
            responded = conn.execute(
                'SELECT COUNT(*) AS c FROM survey_responses WHERE survey_id=?',
                (survey_id,),
            ).fetchone()['c']
            # 응답이 쌓인 뒤 문항을 갈아엎으면 기존 답변이 어느 질문 것인지
            # 알 수 없게 되므로 제목·기간·문자내용만 바꾸도록 막는다.
            existing = _questions_of(conn, survey_id)
            if responded and not _same_questions(existing, questions):
                return _json_error(
                    f'이미 {responded}건의 응답이 있어 문항은 수정할 수 없습니다. '
                    '문항을 바꾸려면 설문을 복제하거나 새로 만들어 주세요.'
                )
            conn.execute(
                '''UPDATE surveys
                      SET title=?, description=?, starts_on=?, ends_on=?,
                          allow_public_link=?, allow_anonymous=?, sms_template=?,
                          updated_at=CURRENT_TIMESTAMP
                    WHERE id=?''',
                (title, description, starts_on, ends_on,
                 allow_public_link, allow_anonymous, template, survey_id),
            )
            if not responded:
                conn.execute('DELETE FROM survey_questions WHERE survey_id=?', (survey_id,))
                _insert_questions(conn, survey_id, questions)
        else:
            cursor = conn.execute(
                '''INSERT INTO surveys
                       (title, description, status, public_token, starts_on, ends_on,
                        allow_public_link, allow_anonymous, sms_template,
                        created_by, created_by_name)
                   VALUES (?, ?, 'draft', ?, ?, ?, ?, ?, ?, ?, ?)''',
                (title, description, _new_token(), starts_on, ends_on,
                 allow_public_link, allow_anonymous, template,
                 actor['emp_no'], actor['name']),
            )
            survey_id = cursor.lastrowid
            _insert_questions(conn, survey_id, questions)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return jsonify({'status': 'success', 'id': survey_id})


def _insert_questions(conn, survey_id, questions):
    conn.executemany(
        '''INSERT INTO survey_questions
               (survey_id, sort_order, qtype, title, options, is_required)
           VALUES (?, ?, ?, ?, ?, ?)''',
        [(survey_id, q['sort_order'], q['qtype'], q['title'], q['options'], q['is_required'])
         for q in questions],
    )


def _same_questions(existing, incoming):
    if len(existing) != len(incoming):
        return False
    for before, after in zip(existing, incoming):
        if before['qtype'] != after['qtype'] or before['title'] != after['title']:
            return False
        if before['options'] != _load_options(after['options']):
            return False
        if int(before['is_required']) != int(after['is_required']):
            return False
    return True


@survey_bp.route('/api/<int:survey_id>/status', methods=['POST'])
def api_status(survey_id):
    payload = request.get_json(silent=True) or {}
    status = str(payload.get('status') or '').strip()
    if status not in SURVEY_STATUSES:
        return _json_error('설문 상태 값을 확인해 주세요.')
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        if not _can_manage(row):
            return _json_error('이 설문을 변경할 권한이 없습니다.', 403)
        conn.execute(
            'UPDATE surveys SET status=?, updated_at=CURRENT_TIMESTAMP WHERE id=?',
            (status, survey_id),
        )
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success', 'survey_status': status})


@survey_bp.route('/api/<int:survey_id>/delete', methods=['POST'])
def api_delete(survey_id):
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        if not _can_manage(row):
            return _json_error('이 설문을 삭제할 권한이 없습니다.', 403)
        conn.execute('DELETE FROM surveys WHERE id=?', (survey_id,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'success'})


@survey_bp.route('/api/<int:survey_id>/duplicate', methods=['POST'])
def api_duplicate(survey_id):
    actor = _actor()
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        cursor = conn.execute(
            '''INSERT INTO surveys
                   (title, description, status, public_token, starts_on, ends_on,
                    allow_public_link, allow_anonymous, sms_template,
                    created_by, created_by_name)
               VALUES (?, ?, 'draft', ?, ?, ?, ?, ?, ?, ?, ?)''',
            (f"{row['title']} (복사본)"[:200], row['description'], _new_token(),
             row['starts_on'], row['ends_on'], row['allow_public_link'],
             row['allow_anonymous'], row['sms_template'],
             actor['emp_no'], actor['name']),
        )
        new_id = cursor.lastrowid
        conn.execute(
            '''INSERT INTO survey_questions
                   (survey_id, sort_order, qtype, title, options, is_required)
               SELECT ?, sort_order, qtype, title, options, is_required
                 FROM survey_questions WHERE survey_id=?''',
            (new_id, survey_id),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return jsonify({'status': 'success', 'id': new_id})


# ---------------------------------------------------------------------------
# 수신자 목록
# ---------------------------------------------------------------------------
def _parse_recipient_lines(raw):
    """'이름,010-0000-0000' 또는 번호만 있는 여러 줄 입력을 해석한다."""
    parsed, errors, seen = [], [], set()
    for index, line in enumerate(str(raw or '').replace(';', '\n').splitlines(), start=1):
        text = line.strip()
        if not text:
            continue
        parts = [part.strip() for part in re.split(r'[,\t|]', text) if part.strip()]
        name, phone_source = '', parts[0] if parts else ''
        if len(parts) >= 2:
            # 앞이 번호면 뒤가 이름, 아니면 앞이 이름.
            if re.search(r'\d', parts[0]) and not re.search(r'\d', parts[1]):
                name, phone_source = parts[1], parts[0]
            else:
                name, phone_source = parts[0], parts[1]
        try:
            phone = normalize_phone(phone_source, required=True)
        except ValueError as exc:
            errors.append(f'{index}번째 줄({text}): {exc}')
            continue
        if phone in seen:
            continue
        seen.add(phone)
        parsed.append({'name': name[:40], 'phone': phone, 'category': '', 'note': ''})
    return parsed, errors


def build_recipient_workbook(include_samples=True):
    """수신자 일괄등록 엑셀 양식을 만든다. 저장해 둔 파일 없이 요청할 때마다 새로 만든다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '수신자목록'
    sheet.append(list(RECIPIENT_COLUMNS))
    header_fill = PatternFill('solid', fgColor='004EA2')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.freeze_panes = 'A2'
    if include_samples:
        for row in RECIPIENT_SAMPLE_ROWS:
            sheet.append(list(row))
    for index, width in enumerate((8, 16, 14, 20, 30), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    # 휴대폰번호 앞의 0이 사라지지 않도록 열 서식을 문자열로 지정한다.
    # (셀을 미리 순회하면 빈 행까지 사용 영역으로 잡히므로 열 단위로 지정한다.)
    phone_column = sheet.column_dimensions['D']
    phone_column.number_format = '@'
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = '@'

    guide = workbook.create_sheet('작성안내')
    for line in (
        ('항목', '설명'),
        ('번호', '보기 편하도록 매기는 순번입니다. 등록할 때는 사용하지 않습니다.'),
        ('구분', '학교명·부서 등 원하는 분류를 적습니다. 수신자 목록에 그대로 표시됩니다.'),
        ('이름', '수신자 이름입니다. 비워 두어도 등록됩니다.'),
        ('핸드폰번호', '실제 문자 발송에 사용합니다. 010-1234-5678 또는 01012345678 형식.'),
        ('비고', '메모입니다. 수신자 목록에 그대로 표시되며 문자에는 들어가지 않습니다.'),
        ('', ''),
        ('중복', '이미 등록된 번호는 자동으로 건너뜁니다.'),
        ('형식오류', '휴대폰번호 형식이 아니면 해당 줄만 건너뛰고 사유를 알려 줍니다.'),
    ):
        guide.append(list(line))
    for cell in guide[1]:
        cell.font = Font(bold=True)
    guide.column_dimensions['A'].width = 14
    guide.column_dimensions['B'].width = 68
    return workbook


@survey_bp.route('/recipients/template')
def recipients_template():
    workbook = build_recipient_workbook(
        include_samples=request.args.get('samples', '1') != '0'
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='설문조사_수신자목록_양식.xlsx',
    )


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_recipient_workbook(stream):
    """업로드된 엑셀에서 (구분, 이름, 핸드폰번호, 비고)를 읽는다.

    머리글 행을 찾아 열 위치를 이름으로 잡으므로 열 순서가 조금 달라도 읽어낸다.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(stream, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError('엑셀 파일을 열 수 없습니다. xlsx 형식인지 확인해 주세요.') from exc

    sheet = workbook['수신자목록'] if '수신자목록' in workbook.sheetnames else workbook.worksheets[0]
    rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not rows:
        raise ValueError('엑셀에 내용이 없습니다.')

    header_index, columns = None, {}
    for index, row in enumerate(rows[:20]):
        labels = {value.replace(' ', ''): position for position, value in enumerate(row) if value}
        if '핸드폰번호' in labels or '휴대폰번호' in labels or '연락처' in labels:
            columns = {
                'category': labels.get('구분'),
                'name': labels.get('이름') if labels.get('이름') is not None else labels.get('성명'),
                'phone': (labels.get('핸드폰번호') if labels.get('핸드폰번호') is not None
                          else labels.get('휴대폰번호') if labels.get('휴대폰번호') is not None
                          else labels.get('연락처')),
                'note': labels.get('비고'),
            }
            header_index = index
            break
    if header_index is None:
        raise ValueError(
            '머리글에서 "핸드폰번호" 열을 찾지 못했습니다. 양식을 내려받아 그대로 사용해 주세요.'
        )

    def pick(row, key):
        position = columns.get(key)
        if position is None or position >= len(row):
            return ''
        return row[position]

    parsed, errors, seen = [], [], set()
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(row):
            continue
        raw_phone = pick(row, 'phone')
        name = pick(row, 'name')
        if not raw_phone and not name:
            continue
        try:
            phone = normalize_phone(raw_phone, required=True)
        except ValueError as exc:
            errors.append(f'{offset}행({name or raw_phone or "빈 줄"}): {exc}')
            continue
        if phone in seen:
            continue
        seen.add(phone)
        parsed.append({
            'name': name[:40],
            'phone': phone,
            'category': pick(row, 'category')[:40],
            'note': pick(row, 'note')[:200],
        })
    return parsed, errors


def _store_recipients(conn, survey_id, parsed):
    added, skipped = 0, 0
    for item in parsed:
        exists = conn.execute(
            'SELECT id FROM survey_recipients WHERE survey_id=? AND phone=?',
            (survey_id, item['phone']),
        ).fetchone()
        if exists:
            skipped += 1
            continue
        conn.execute(
            '''INSERT INTO survey_recipients (survey_id, name, phone, category, note, token)
               VALUES (?, ?, ?, ?, ?, ?)''',
            (survey_id, item['name'], item['phone'], item.get('category') or '',
             item.get('note') or '', _new_token()),
        )
        added += 1
    return added, skipped


@survey_bp.route('/api/<int:survey_id>/recipients/upload', methods=['POST'])
def api_upload_recipients(survey_id):
    upload = request.files.get('file')
    if not upload or not upload.filename:
        return _json_error('업로드할 엑셀 파일을 선택해 주세요.')
    if not upload.filename.lower().endswith(('.xlsx', '.xlsm')):
        return _json_error('엑셀(xlsx) 파일만 등록할 수 있습니다.')
    try:
        parsed, errors = parse_recipient_workbook(upload.stream)
    except ValueError as exc:
        return _json_error(str(exc))
    if not parsed:
        return _json_error(
            '등록할 휴대폰번호가 없습니다.' + (' ' + errors[0] if errors else '')
        )

    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        if not _can_manage(row):
            return _json_error('이 설문의 수신자를 편집할 권한이 없습니다.', 403)
        current = conn.execute(
            'SELECT COUNT(*) AS c FROM survey_recipients WHERE survey_id=?', (survey_id,)
        ).fetchone()['c']
        if current + len(parsed) > MAX_RECIPIENTS:
            return _json_error(f'수신자는 설문당 최대 {MAX_RECIPIENTS}명까지 등록할 수 있습니다.')
        added, skipped = _store_recipients(conn, survey_id, parsed)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return jsonify({
        'status': 'success',
        'added': added,
        'skipped': skipped,
        'errors': errors[:20],
    })


@survey_bp.route('/api/<int:survey_id>/recipients', methods=['POST'])
def api_add_recipients(survey_id):
    payload = request.get_json(silent=True) or {}
    parsed, errors = _parse_recipient_lines(payload.get('raw'))
    if not parsed:
        return _json_error(
            '등록할 휴대폰번호가 없습니다.' + (' ' + errors[0] if errors else '')
        )
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        if not _can_manage(row):
            return _json_error('이 설문의 수신자를 편집할 권한이 없습니다.', 403)
        current = conn.execute(
            'SELECT COUNT(*) AS c FROM survey_recipients WHERE survey_id=?', (survey_id,)
        ).fetchone()['c']
        if current + len(parsed) > MAX_RECIPIENTS:
            return _json_error(f'수신자는 설문당 최대 {MAX_RECIPIENTS}명까지 등록할 수 있습니다.')
        added, skipped = _store_recipients(conn, survey_id, parsed)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return jsonify({
        'status': 'success',
        'added': added,
        'skipped': skipped,
        'errors': errors[:20],
    })


@survey_bp.route('/api/<int:survey_id>/recipients/delete', methods=['POST'])
def api_delete_recipients(survey_id):
    payload = request.get_json(silent=True) or {}
    ids = payload.get('ids')
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        if not _can_manage(row):
            return _json_error('이 설문의 수신자를 편집할 권한이 없습니다.', 403)
        if payload.get('all'):
            cursor = conn.execute(
                'DELETE FROM survey_recipients WHERE survey_id=?', (survey_id,)
            )
        else:
            target = [int(value) for value in ids or [] if str(value).isdigit()]
            if not target:
                return _json_error('삭제할 수신자를 선택해 주세요.')
            placeholders = ','.join('?' * len(target))
            cursor = conn.execute(
                f'DELETE FROM survey_recipients WHERE survey_id=? AND id IN ({placeholders})',
                [survey_id, *target],
            )
        conn.commit()
        deleted = cursor.rowcount
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return jsonify({'status': 'success', 'deleted': deleted})


# ---------------------------------------------------------------------------
# 문자 발송
# ---------------------------------------------------------------------------
@survey_bp.route('/api/<int:survey_id>/preview', methods=['POST'])
def api_preview(survey_id):
    payload = request.get_json(silent=True) or {}
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        token = row['public_token']
        title = row['title']
        saved_template = row['sms_template']
    finally:
        conn.close()
    solapi = get_solapi_settings()
    text = _render_sms(
        payload.get('sms_template') or saved_template,
        title=title,
        link=_public_link(token, resolve_public_origin(solapi)),
        sender_name=solapi.get('sender_name') or DEFAULT_SENDER_NAME,
    )
    length = message_byte_length(text)
    return jsonify({
        'status': 'success',
        'text': text,
        'bytes': length,
        'message_type': 'SMS' if length <= SMS_BYTE_LIMIT else 'LMS',
    })


@survey_bp.route('/api/<int:survey_id>/send', methods=['POST'])
def api_send(survey_id):
    payload = request.get_json(silent=True) or {}
    resend = bool(payload.get('resend'))
    ids = [int(value) for value in payload.get('ids') or [] if str(value).isdigit()]

    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        if not _can_manage(row):
            return _json_error('이 설문을 발송할 권한이 없습니다.', 403)
        if not _questions_of(conn, survey_id):
            return _json_error('설문 문항이 없습니다. 문항을 먼저 저장해 주세요.')
        if row['status'] == 'closed':
            return _json_error('마감된 설문은 발송할 수 없습니다.')

        sql = 'SELECT * FROM survey_recipients WHERE survey_id=?'
        params = [survey_id]
        if ids:
            sql += ' AND id IN (' + ','.join('?' * len(ids)) + ')'
            params.extend(ids)
        elif not resend:
            sql += " AND send_status<>'sent'"
        targets = conn.execute(sql + ' ORDER BY id', params).fetchall()
        if not targets:
            return _json_error(
                '발송할 수신자가 없습니다. 이미 모두 발송했다면 재발송을 선택해 주세요.'
            )
        title = row['title']
        template = row['sms_template']
        # 발송 시점에 진행중으로 바꿔 응답 링크가 바로 열리게 한다.
        if row['status'] == 'draft':
            conn.execute(
                "UPDATE surveys SET status='open', updated_at=CURRENT_TIMESTAMP WHERE id=?",
                (survey_id,),
            )
            conn.commit()
    finally:
        conn.close()

    solapi = get_solapi_settings()
    if not solapi.get('sms_configured'):
        return _json_error(
            '통합관리 > 솔라피설정에서 API KEY·API SECRET·발신번호를 먼저 저장해 주세요.'
        )
    sender_name = solapi.get('sender_name') or DEFAULT_SENDER_NAME
    origin = resolve_public_origin(solapi)

    messages = [{
        'to': item['phone'],
        'subject': f'{sender_name} 설문조사',
        'text': _render_sms(
            template,
            title=title,
            link=_public_link(item['token'], origin),
            sender_name=sender_name,
        ),
    } for item in targets]

    try:
        results = send_bulk_text(messages, settings=solapi)
    except (RuntimeError, ValueError) as exc:
        return _json_error(str(exc), 502)

    success = sum(1 for item in results if item['ok'])
    failed = len(results) - success
    actor = _actor()
    conn = get_db()
    try:
        for target, outcome in zip(targets, results):
            conn.execute(
                '''UPDATE survey_recipients
                      SET send_status=?, send_error=?, message_id=?, sent_at=CURRENT_TIMESTAMP
                    WHERE id=?''',
                ('sent' if outcome['ok'] else 'failed',
                 '' if outcome['ok'] else outcome['error'][:300],
                 outcome['message_id'], target['id']),
            )
        conn.execute(
            '''INSERT INTO survey_send_logs (survey_id, sent_by, total, success, failed, memo)
               VALUES (?, ?, ?, ?, ?, ?)''',
            (survey_id, actor['name'] or actor['emp_no'], len(results), success, failed,
             '재발송 포함' if resend else ''),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    return jsonify({
        'status': 'success',
        'total': len(results),
        'success': success,
        'failed': failed,
        'failures': [
            {'phone': mask_phone(item['to']), 'error': item['error']}
            for item in results if not item['ok']
        ][:30],
    })


# ---------------------------------------------------------------------------
# 결과 및 통계
# ---------------------------------------------------------------------------
def _build_statistics(conn, survey_id):
    questions = _questions_of(conn, survey_id)
    answers = conn.execute(
        '''SELECT a.question_id, a.answer_text, a.answer_options
             FROM survey_answers a
             JOIN survey_responses r ON r.id = a.response_id
            WHERE r.survey_id=?''',
        (survey_id,),
    ).fetchall()

    grouped = {question['id']: [] for question in questions}
    for row in answers:
        grouped.setdefault(row['question_id'], []).append(row)

    stats = []
    for question in questions:
        rows = grouped.get(question['id'], [])
        entry = {
            'id': question['id'],
            'title': question['title'],
            'qtype': question['qtype'],
            'qtype_label': question['qtype_label'],
            'answer_count': len(rows),
            'options': [],
            'texts': [],
            'average': None,
        }
        if question['qtype'] == 'text':
            entry['texts'] = [
                str(row['answer_text'] or '').strip()
                for row in rows if str(row['answer_text'] or '').strip()
            ][:500]
            stats.append(entry)
            continue

        counter = {label: 0 for label in question['options']}
        picked_total = 0
        for row in rows:
            for label in _load_options(row['answer_options']):
                if label in counter:
                    counter[label] += 1
                    picked_total += 1
        base = picked_total if question['qtype'] == 'multiple' else len(rows)
        entry['options'] = [{
            'label': label,
            'count': counter[label],
            'ratio': round(counter[label] * 100 / base, 1) if base else 0.0,
        } for label in question['options']]
        if question['qtype'] == 'scale' and question['options']:
            # 첫 보기를 만점으로 두고 5→1점 순으로 환산한 평균 점수.
            size = len(question['options'])
            total_score = sum(
                counter[label] * (size - position)
                for position, label in enumerate(question['options'])
            )
            answered = sum(counter.values())
            entry['average'] = round(total_score / answered, 2) if answered else None
        stats.append(entry)
    return questions, stats


def _answer_display(row):
    """보기 선택은 쉼표로 잇고, 주관식은 입력 원문을 그대로 쓴다."""
    options = _load_options(row['answer_options'])
    return ', '.join(options) if options else str(row['answer_text'] or '')


def _respondent_rows(conn, survey_id, questions, limit=500):
    """누가 어떤 문항에 무엇이라고 답했는지 한 줄씩 만든다.

    무기명 응답은 수신자와 연결하지 않았으므로 이름·연락처 없이 응답 내용만 남는다.
    """
    rows = conn.execute(
        '''SELECT r.id, r.submitted_at, r.is_anonymous, c.name, c.phone, c.category, c.note
             FROM survey_responses r
             LEFT JOIN survey_recipients c ON c.id = r.recipient_id
            WHERE r.survey_id=? ORDER BY r.id DESC LIMIT ?''',
        (survey_id, limit),
    ).fetchall()
    if not rows:
        return []

    answers = {}
    placeholders = ','.join('?' * len(rows))
    for item in conn.execute(
        f'''SELECT response_id, question_id, answer_text, answer_options
              FROM survey_answers WHERE response_id IN ({placeholders})''',
        [row['id'] for row in rows],
    ).fetchall():
        answers[(item['response_id'], item['question_id'])] = _answer_display(item)

    result = []
    for row in rows:
        anonymous = bool(row['is_anonymous'])
        result.append({
            'id': row['id'],
            'anonymous': anonymous,
            'name': '무기명' if anonymous else (row['name'] or '익명'),
            'category': '' if anonymous else (row['category'] or ''),
            'note': '' if anonymous else (row['note'] or ''),
            'phone': '-' if anonymous or not row['phone'] else mask_phone(row['phone']),
            'submitted_at': str(row['submitted_at'] or '')[:16],
            'answers': [answers.get((row['id'], question['id']), '')
                        for question in questions],
        })
    return result


@survey_bp.route('/api/<int:survey_id>/stats')
def api_stats(survey_id):
    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return _json_error('설문을 찾을 수 없습니다.', 404)
        questions, stats = _build_statistics(conn, survey_id)
        summary = conn.execute(
            '''SELECT
                 (SELECT COUNT(*) FROM survey_responses WHERE survey_id=?) AS responses,
                 (SELECT COUNT(*) FROM survey_recipients WHERE survey_id=?) AS recipients,
                 (SELECT COUNT(*) FROM survey_recipients
                   WHERE survey_id=? AND send_status='sent') AS sent,
                 (SELECT COUNT(*) FROM survey_recipients
                   WHERE survey_id=? AND responded_at IS NOT NULL) AS responded''',
            (survey_id, survey_id, survey_id, survey_id),
        ).fetchone()
        responses = _respondent_rows(conn, survey_id, questions)
        survey = _survey_dict(
            row,
            response_count=summary['responses'],
            recipient_count=summary['recipients'],
            sent_count=summary['sent'],
        )
    finally:
        conn.close()
    # 문자를 보내지 않고 링크만 공유한 설문도 있으므로, 발송 건수가 없으면
    # 등록 수신자 수를 분모로 삼는다.
    base = summary['sent'] or summary['recipients'] or 0
    return jsonify({
        'status': 'success',
        'survey': survey,
        'questions': questions,
        'stats': stats,
        'responses': responses,
        'summary': {
            'responses': summary['responses'],
            'recipients': summary['recipients'],
            'sent': summary['sent'],
            'responded': summary['responded'],
            'rate': round(summary['responded'] * 100 / base, 1) if base else 0.0,
        },
    })


@survey_bp.route('/<int:survey_id>/export')
def export_result(survey_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    conn = get_db()
    try:
        row = _survey_row(conn, survey_id)
        if not row:
            return '설문을 찾을 수 없습니다.', 404
        questions, stats = _build_statistics(conn, survey_id)
        respondents = _respondent_rows(conn, survey_id, questions, limit=100000)
        recipients = conn.execute(
            '''SELECT name, category, phone, note, send_status, sent_at, responded_at
                 FROM survey_recipients WHERE survey_id=? ORDER BY id''',
            (survey_id,),
        ).fetchall()
    finally:
        conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '응답자별'
    header = ['번호', '구분', '이름', '휴대폰', '응답일시'] + [
        f"{index}. {question['title']}" for index, question in enumerate(questions, start=1)
    ]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'
    for number, respondent in enumerate(reversed(respondents), start=1):
        sheet.append([
            number,
            respondent['category'],
            respondent['name'],
            respondent['phone'],
            respondent['submitted_at'],
            *respondent['answers'],
        ])

    recipient_sheet = workbook.create_sheet('수신자')
    recipient_sheet.append(['번호', '구분', '이름', '핸드폰번호', '비고',
                            '발송상태', '발송일시', '응답일시'])
    for cell in recipient_sheet[1]:
        cell.font = Font(bold=True)
    for number, item in enumerate(recipients, start=1):
        recipient_sheet.append([
            number,
            item['category'] or '',
            item['name'] or '',
            format_phone(item['phone']),
            item['note'] or '',
            {'ready': '대기', 'sent': '발송', 'failed': '실패'}.get(
                item['send_status'], item['send_status']),
            str(item['sent_at'] or '')[:16],
            str(item['responded_at'] or '')[:16] or '미응답',
        ])

    summary_sheet = workbook.create_sheet('통계')
    summary_sheet.append(['문항', '유형', '보기', '응답수', '비율(%)'])
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True)
    for entry in stats:
        if entry['qtype'] == 'text':
            summary_sheet.append([entry['title'], entry['qtype_label'], '주관식 응답',
                                  entry['answer_count'], ''])
            continue
        for option in entry['options']:
            summary_sheet.append([entry['title'], entry['qtype_label'], option['label'],
                                  option['count'], option['ratio']])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    safe_title = re.sub(r'[\\/:*?"<>|]', '_', str(row['title']))[:60] or '설문결과'
    return send_file(
        buffer,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f"{safe_title}_설문결과.xlsx",
    )


# ---------------------------------------------------------------------------
# 공개 응답 화면 (로그인 없음)
# ---------------------------------------------------------------------------
def _resolve_token(conn, token):
    """수신자 개인 링크와 공용 링크를 모두 받아 (설문, 수신자)를 돌려준다."""
    text = str(token or '').strip()
    if not text:
        return None, None
    recipient = conn.execute(
        'SELECT * FROM survey_recipients WHERE token=?', (text,)
    ).fetchone()
    if recipient:
        return _survey_row(conn, recipient['survey_id']), recipient
    survey = conn.execute(
        'SELECT * FROM surveys WHERE public_token=?', (text,)
    ).fetchone()
    if survey and survey['allow_public_link']:
        return survey, None
    return survey, None


@survey_bp.route('/r/<token>')
def respond(token):
    conn = get_db()
    try:
        survey, recipient = _resolve_token(conn, token)
        if not survey:
            return render_template(
                'survey/public.html', survey=None,
                notice='설문 주소가 올바르지 않거나 삭제된 설문입니다.',
            ), 404
        if recipient is None and not survey['allow_public_link']:
            return render_template(
                'survey/public.html', survey=None,
                notice='개인별로 발송된 설문 주소로만 참여할 수 있습니다.',
            ), 403
        open_now, notice = _survey_is_open(survey)
        questions = _questions_of(conn, survey['id'])
        already = False
        if recipient is not None and recipient['responded_at']:
            already = True
    finally:
        conn.close()

    if already:
        return render_template(
            'survey/public.html', survey=survey, questions=[],
            notice='이미 응답을 완료하셨습니다. 참여해 주셔서 감사합니다.',
            done=True,
        )
    if not open_now:
        return render_template(
            'survey/public.html', survey=survey, questions=[], notice=notice,
        )
    return render_template(
        'survey/public.html',
        survey=survey,
        questions=questions,
        recipient_name=(recipient['name'] if recipient is not None else ''),
        # 공용 링크로 들어온 사람은 애초에 신원이 없으므로 선택지를 보여주지 않는다.
        allow_anonymous=bool(survey['allow_anonymous'] and recipient is not None),
        notice='',
        token=token,
    )


@survey_bp.route('/r/<token>', methods=['POST'])
def submit_response(token):
    payload = request.get_json(silent=True) or {}
    answers = payload.get('answers')
    if not isinstance(answers, dict):
        return _json_error('응답 형식을 확인해 주세요.')

    conn = get_db()
    try:
        survey, recipient = _resolve_token(conn, token)
        if not survey:
            return _json_error('설문 주소가 올바르지 않습니다.', 404)
        if recipient is None and not survey['allow_public_link']:
            return _json_error('개인별로 발송된 설문 주소로만 참여할 수 있습니다.', 403)
        open_now, notice = _survey_is_open(survey)
        if not open_now:
            return _json_error(notice, 403)
        if recipient is not None and recipient['responded_at']:
            return _json_error('이미 응답을 완료하셨습니다.', 409)
        # 무기명을 고르면 응답에 수신자를 연결하지 않는다. 다만 중복 응답을 막고
        # 응답률을 세기 위해 "응답했다"는 사실만 수신자 쪽에 남긴다.
        anonymous = bool(survey['allow_anonymous']) and bool(payload.get('anonymous'))

        questions = _questions_of(conn, survey['id'])
        prepared = []
        for question in questions:
            raw = answers.get(str(question['id']))
            if question['qtype'] == 'text':
                text = str(raw or '').strip()[:2000]
                if question['is_required'] and not text:
                    return _json_error(f"'{question['title']}' 문항에 답변해 주세요.")
                prepared.append((question['id'], text, json.dumps([], ensure_ascii=False)))
                continue
            picked = raw if isinstance(raw, list) else ([raw] if raw else [])
            picked = [str(item) for item in picked if str(item) in question['options']]
            if question['qtype'] in {'single', 'scale'}:
                picked = picked[:1]
            if question['is_required'] and not picked:
                return _json_error(f"'{question['title']}' 문항을 선택해 주세요.")
            prepared.append((
                question['id'], '', json.dumps(picked, ensure_ascii=False),
            ))

        cursor = conn.execute(
            '''INSERT INTO survey_responses (survey_id, recipient_id, is_anonymous)
               VALUES (?, ?, ?)''',
            (survey['id'],
             None if anonymous or recipient is None else recipient['id'],
             1 if anonymous else 0),
        )
        response_id = cursor.lastrowid
        conn.executemany(
            '''INSERT INTO survey_answers (response_id, question_id, answer_text, answer_options)
               VALUES (?, ?, ?, ?)''',
            [(response_id, question_id, text, options)
             for question_id, text, options in prepared],
        )
        if recipient is not None:
            conn.execute(
                'UPDATE survey_recipients SET responded_at=CURRENT_TIMESTAMP WHERE id=?',
                (recipient['id'],),
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return jsonify({'status': 'success'})
